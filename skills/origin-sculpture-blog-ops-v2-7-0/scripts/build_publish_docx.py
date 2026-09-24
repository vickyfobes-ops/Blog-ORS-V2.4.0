#!/usr/bin/env python3
"""Build a clean, site-styled Word copy of an Origin Sculpture article.

The document mirrors the current Origin Sculpture article typography while
remaining a portable editorial handoff: featured image, public title, article
body, contextual links, tables, lists, and inline images only. It deliberately
omits review labels, metadata tables, sources, headers, footers, and page
numbers.
"""

from __future__ import annotations

import argparse
import json
import re
import struct
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree as ET


BLACK = "000000"
WHITE = "FFFFFF"
TABLE_HEADER = "F2F2F2"
TABLE_RULE = "BFBFBF"
BODY_FONT = "Poppins"
HEADING_FONT = "Libre Baskerville"
BODY_STYLE = "Origin Body"
TITLE_STYLE = "Origin Title"
H2_STYLE = "Origin H2"
H3_STYLE = "Origin H3"
BULLET_STYLE = "Origin Bullet"
NUMBER_STYLE = "Origin Number"
TABLE_STYLE = "Origin Publish Table"


def set_font_name(font, name: str) -> None:
    font.name = name
    r_fonts = font._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in r_fonts.attrib:
            del r_fonts.attrib[key]
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)


def paragraph_style(document: Document, name: str, base: str):
    """Return a deterministic custom paragraph style, creating it once."""
    try:
        style = document.styles[name]
    except KeyError:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles[base]
    return style


def set_style_tracking(style, twentieths_of_a_point: int) -> None:
    r_pr = style._element.get_or_add_rPr()
    spacing = r_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        r_pr.append(spacing)
    spacing.set(qn("w:val"), str(twentieths_of_a_point))


def clear_story(story) -> None:
    for paragraph in story.paragraphs:
        for run in list(paragraph.runs):
            paragraph._p.remove(run._r)


def set_cell_margins(cell, *, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), TABLE_RULE)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.90)
    section.right_margin = Inches(0.90)
    section.header_distance = Inches(0.10)
    section.footer_distance = Inches(0.10)
    clear_story(section.header)
    clear_story(section.footer)

    normal = document.styles["Normal"]
    set_font_name(normal.font, BODY_FONT)
    normal.font.size = Pt(12.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    body = paragraph_style(document, BODY_STYLE, "Normal")
    set_font_name(body.font, BODY_FONT)
    body.font.size = Pt(12.5)  # Stable half-point approximation of site body type.
    body.font.color.rgb = RGBColor.from_string(BLACK)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(15)  # Site paragraph margin: 20 CSS px.
    body.paragraph_format.line_spacing = 1.7
    body.paragraph_format.widow_control = True
    set_style_tracking(body, -10)  # Site letter-spacing: -0.04em.

    title = paragraph_style(document, TITLE_STYLE, BODY_STYLE)
    set_font_name(title.font, HEADING_FONT)
    title.font.size = Pt(33)  # Site title: 44 CSS px.
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(30)
    title.paragraph_format.line_spacing = 1.2
    title.paragraph_format.keep_with_next = True
    set_style_tracking(title, -26)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    h2 = paragraph_style(document, H2_STYLE, BODY_STYLE)
    set_font_name(h2.font, HEADING_FONT)
    h2.font.size = Pt(27)  # Site H2: 36 CSS px.
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLACK)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(30)
    h2.paragraph_format.space_after = Pt(15)
    h2.paragraph_format.line_spacing = 1.0
    h2.paragraph_format.keep_with_next = True
    set_style_tracking(h2, -22)

    h3 = paragraph_style(document, H3_STYLE, BODY_STYLE)
    set_font_name(h3.font, HEADING_FONT)
    h3.font.size = Pt(16.5)  # Site H3: 22 CSS px.
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(BLACK)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(24)
    h3.paragraph_format.space_after = Pt(12)
    h3.paragraph_format.line_spacing = 1.0
    h3.paragraph_format.keep_with_next = True
    set_style_tracking(h3, -13)

    for style_name, base_style in (
        (BULLET_STYLE, "List Bullet"),
        (NUMBER_STYLE, BODY_STYLE),
    ):
        style = paragraph_style(document, style_name, base_style)
        set_font_name(style.font, BODY_FONT)
        style.font.size = Pt(12.5)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.left_indent = Inches(0.30)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.7
        set_style_tracking(style, -10)

    table_text = paragraph_style(document, TABLE_STYLE, BODY_STYLE)
    set_font_name(table_text.font, BODY_FONT)
    table_text.font.size = Pt(10.5)
    table_text.font.color.rgb = RGBColor.from_string(BLACK)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.35
    set_style_tracking(table_text, -8)

def configure_numbered_list(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "520")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "520")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), BODY_FONT)
    r_pr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLACK)
    r_pr.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "25")
    r_pr.append(size)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(num_id_node)


def font_embedding_allowed(font_path: Path) -> bool:
    data = font_path.read_bytes()
    if len(data) < 12:
        return False
    num_tables = struct.unpack(">H", data[4:6])[0]
    for index in range(num_tables):
        start = 12 + index * 16
        tag, _checksum, offset, length = struct.unpack(">4sIII", data[start : start + 16])
        if tag != b"OS/2" or length < 10:
            continue
        fs_type = struct.unpack(">H", data[offset + 8 : offset + 10])[0]
        return not bool(fs_type & 0x0002)
    return True


def obfuscate_font(font_data: bytes, font_key: uuid.UUID) -> bytes:
    data = bytearray(font_data)
    key = bytes.fromhex(font_key.hex)[::-1]
    for index in range(min(32, len(data))):
        data[index] ^= key[index % 16]
    return bytes(data)


def embed_font_family(
    document_path: Path,
    family_name: str,
    regular_path: Path,
    bold_path: Path,
    seed_prefix: str,
) -> None:
    """Embed one OFL-licensed regular/bold font family in the DOCX."""
    for font_path in (regular_path, bold_path):
        if not font_path.exists():
            raise FileNotFoundError(f"Required font asset is missing: {font_path}")
        if not font_embedding_allowed(font_path):
            raise ValueError(f"Font embedding is restricted by fsType: {font_path}")

    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"

    with zipfile.ZipFile(document_path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}

    font_table = ET.fromstring(entries["word/fontTable.xml"])
    font_node = None
    for candidate in font_table.findall(f"{{{w_ns}}}font"):
        if candidate.get(f"{{{w_ns}}}name") == family_name:
            font_node = candidate
            break
    if font_node is None:
        font_node = ET.SubElement(font_table, f"{{{w_ns}}}font", {f"{{{w_ns}}}name": family_name})
    for tag in ("embedRegular", "embedBold"):
        existing = font_node.find(f"{{{w_ns}}}{tag}")
        if existing is not None:
            font_node.remove(existing)

    rels_name = "word/_rels/fontTable.xml.rels"
    if rels_name in entries:
        rels = ET.fromstring(entries[rels_name])
    else:
        rels = ET.Element(f"{{{rel_ns}}}Relationships", nsmap={None: rel_ns})
    existing_ids = {
        rel.get("Id", "")
        for rel in rels.findall(f"{{{rel_ns}}}Relationship")
    }
    next_id = 1
    while f"rId{next_id}" in existing_ids:
        next_id += 1

    font_parts: dict[str, bytes] = {}
    for tag, font_path, seed in (
        ("embedRegular", regular_path, f"{seed_prefix}-regular"),
        ("embedBold", bold_path, f"{seed_prefix}-bold"),
    ):
        rel_id = f"rId{next_id}"
        next_id += 1
        font_key = uuid.uuid5(uuid.NAMESPACE_URL, seed)
        part_name = f"{str(font_key).upper()}.odttf"
        ET.SubElement(
            rels,
            f"{{{rel_ns}}}Relationship",
            {
                "Id": rel_id,
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
                "Target": f"fonts/{part_name}",
            },
        )
        ET.SubElement(
            font_node,
            f"{{{w_ns}}}{tag}",
            {
                f"{{{r_ns}}}id": rel_id,
                f"{{{w_ns}}}fontKey": "{" + str(font_key).upper() + "}",
            },
        )
        font_parts[f"word/fonts/{part_name}"] = obfuscate_font(font_path.read_bytes(), font_key)

    content_types = ET.fromstring(entries["[Content_Types].xml"])
    has_odttf = any(
        node.get("Extension") == "odttf"
        for node in content_types.findall(f"{{{ct_ns}}}Default")
    )
    if not has_odttf:
        ET.SubElement(
            content_types,
            f"{{{ct_ns}}}Default",
            {
                "Extension": "odttf",
                "ContentType": "application/vnd.openxmlformats-officedocument.obfuscatedFont",
            },
        )

    settings = ET.fromstring(entries["word/settings.xml"])
    if settings.find(f"{{{w_ns}}}embedTrueTypeFonts") is None:
        settings.append(ET.Element(f"{{{w_ns}}}embedTrueTypeFonts"))

    entries["word/fontTable.xml"] = ET.tostring(font_table, encoding="utf-8", xml_declaration=True)
    entries[rels_name] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)
    entries["[Content_Types].xml"] = ET.tostring(content_types, encoding="utf-8", xml_declaration=True)
    entries["word/settings.xml"] = ET.tostring(settings, encoding="utf-8", xml_declaration=True)
    entries.update(font_parts)

    temp_path = document_path.with_suffix(".embedding.docx")
    with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, data in entries.items():
            target.writestr(name, data)
    temp_path.replace(document_path)


def add_hyperlink(paragraph, label: str, url: str, *, bold: bool = False) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_fonts.set(qn("w:eastAsia"), BODY_FONT)
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLACK)
    r_pr.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "25")
    r_pr.append(size)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "-10")
    r_pr.append(spacing)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)")


def add_inline(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            inner = token[2:-2]
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", inner)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2), bold=True)
            else:
                paragraph.add_run(inner).bold = True
        elif token.startswith("["):
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        else:
            run = paragraph.add_run(token[1:-1])
            set_font_name(run.font, BODY_FONT)
            run.font.size = Pt(12.5)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [split_table_row(line) for line in lines]
    if len(rows) > 1 and is_table_separator(lines[1]):
        rows.pop(1)
    columns = len(rows[0])
    if columns == 4:
        widths = [1650, 2000, 2200, 3798]
    else:
        widths = [9648 // columns] * columns

    # LibreOffice can lose or clip a repeated header when a tall four-column
    # comparison table leaves only one data row for the next page. Split that
    # layout deterministically into continuation tables with the header copied
    # into each chunk. Shopify HTML remains unchanged; this is DOCX-only QA.
    data_rows = rows[1:]
    chunks = [data_rows]
    if columns == 4 and len(data_rows) > 4:
        chunks = [data_rows[index : index + 4] for index in range(0, len(data_rows), 4)]

    for chunk_idx, chunk in enumerate(chunks):
        if chunk_idx:
            document.add_page_break()
        table_rows = [rows[0], *chunk]
        table = document.add_table(rows=len(table_rows), cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        for row_idx, values in enumerate(table_rows):
            row = table.rows[row_idx]
            prevent_row_split(row)
            if row_idx == 0:
                set_repeat_table_header(row)
            for col_idx, value in enumerate(values):
                cell = row.cells[col_idx]
                set_cell_width(cell, widths[col_idx])
                set_cell_margins(cell)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_shading(cell, TABLE_HEADER if row_idx == 0 else WHITE)
                paragraph = cell.paragraphs[0]
                paragraph.style = TABLE_STYLE
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline(paragraph, value)
                if row_idx == 0:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string(BLACK)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_picture(
    document: Document,
    image_path: Path,
    alt: str,
    *,
    width: float,
    keep_with_next: bool = False,
) -> None:
    paragraph = document.add_paragraph(style=BODY_STYLE)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(15)
    paragraph.paragraph_format.space_after = Pt(15)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_with_next
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt)


def extract_article_lines(markdown: str) -> list[str]:
    marker = "## Article"
    if marker not in markdown:
        raise ValueError("Article marker not found")
    article = markdown.split(marker, 1)[1].strip()
    lines = article.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    return lines


def normalized_heading(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def parse_image_placements(assets: dict[str, dict]) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    before: dict[str, list[str]] = {}
    after: dict[str, list[str]] = {}
    for slot, asset in assets.items():
        if slot == "cover":
            continue
        placement = str(asset.get("placement", "")).strip()
        match = re.fullmatch(r"(Before|After)\s+(.+)", placement, re.I)
        if not match:
            raise ValueError(
                f"Image slot '{slot}' needs placement 'Before <exact H2/H3>' or 'After <exact H2/H3>'"
            )
        position = match.group(1).lower()
        heading = normalized_heading(match.group(2))
        target = before if position == "before" else after
        target.setdefault(heading, []).append(slot)
    return before, after


def asset_image_path(bundle_dir: Path, asset: dict) -> Path:
    relative = str(asset.get("png", "")).strip()
    if not relative:
        raise ValueError(f"Image slot '{asset.get('slot', '')}' is missing its DOCX PNG asset")
    image_path = (bundle_dir / relative).resolve()
    if bundle_dir != image_path and bundle_dir not in image_path.parents:
        raise ValueError(f"Image path escapes bundle: {relative}")
    if not image_path.is_file() or image_path.suffix.lower() != ".png":
        raise ValueError(f"DOCX image is missing or is not PNG: {relative}")
    return image_path


def add_asset_picture(document: Document, bundle_dir: Path, asset: dict) -> None:
    # Match the latest accepted Origin handoff: body scenes are 5.5 in, while
    # the final project-review image is intentionally 4.5 in so the closing
    # CTA does not spill onto a nearly empty extra page.
    default_width = 4.5 if asset.get("slot") == "project-review" else 5.5
    width = float(asset.get("docxWidthInches", default_width))
    if not 2.0 <= width <= 6.8:
        raise ValueError(f"docxWidthInches must be between 2.0 and 6.8 for slot: {asset.get('slot', '')}")
    add_picture(
        document,
        asset_image_path(bundle_dir, asset),
        str(asset["alt"]),
        width=width,
        keep_with_next=asset.get("slot") == "project-review",
    )


def add_article(
    document: Document,
    lines: list[str],
    bundle_dir: Path,
    assets: dict[str, dict],
    number_id: int,
) -> None:
    before_heading, after_heading = parse_image_placements(assets)
    closing_headings = {
        normalized_heading(str(asset.get("placement", "")).split(None, 1)[1])
        for asset in assets.values()
        if asset.get("slot") == "project-review"
        and str(asset.get("placement", "")).lower().startswith("before ")
    }
    pending_after: list[str] = []
    inserted: set[str] = set()
    closing_section = False

    def insert_slots(slots: list[str]) -> None:
        for slot in slots:
            if slot in inserted:
                continue
            add_asset_picture(document, bundle_dir, assets[slot])
            inserted.add(slot)

    def flush_pending_after() -> None:
        nonlocal pending_after
        insert_slots(pending_after)
        pending_after = []

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_markdown_table(document, table_lines)
            flush_pending_after()
            continue
        if line.startswith("## "):
            flush_pending_after()
            heading = normalized_heading(line[3:])
            closing_section = heading in closing_headings
            insert_slots(before_heading.get(heading, []))
            document.add_paragraph(heading, style=H2_STYLE)
            pending_after = list(after_heading.get(heading, []))
        elif line.startswith("### "):
            flush_pending_after()
            heading = normalized_heading(line[4:])
            closing_section = heading in closing_headings
            insert_slots(before_heading.get(heading, []))
            document.add_paragraph(heading, style=H3_STYLE)
            pending_after = list(after_heading.get(heading, []))
        elif re.match(r"^- ", line):
            paragraph = document.add_paragraph(style=BULLET_STYLE)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = closing_section
            add_inline(paragraph, re.sub(r"^- ", "", line))
            flush_pending_after()
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(style=NUMBER_STYLE)
            paragraph.paragraph_format.left_indent = Inches(0.36)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.7
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = closing_section
            apply_numbering(paragraph, number_id)
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
            flush_pending_after()
        else:
            paragraph = document.add_paragraph(style=BODY_STYLE)
            paragraph.paragraph_format.keep_with_next = closing_section
            add_inline(paragraph, line)
            flush_pending_after()
        i += 1
    flush_pending_after()
    missing = sorted(set(assets) - {"cover"} - inserted)
    if missing:
        raise ValueError(f"Image placements did not match article headings: {', '.join(missing)}")


def build(bundle_dir: Path, output: Path) -> None:
    article_text = (bundle_dir / "article.md").read_text(encoding="utf-8")
    meta = json.loads((bundle_dir / "meta.json").read_text(encoding="utf-8"))
    assets = {
        item["slot"]: item
        for item in json.loads((bundle_dir / "image-assets.json").read_text(encoding="utf-8"))
    }

    document = Document()
    configure_document(document)
    document.core_properties.title = str(meta["title"])
    document.core_properties.author = str(meta.get("author", "Origin Sculpture"))
    document.core_properties.subject = "Final Shopify article copy"
    keywords = [str(meta.get("seoTitle", "")).strip(), *[str(tag).strip() for tag in meta.get("tags", [])]]
    document.core_properties.keywords = ", ".join(dict.fromkeys(value for value in keywords if value))
    document.core_properties.comments = "Clean publication copy; no review headers, footers, or notes."
    number_id = configure_numbered_list(document)

    cover = assets["cover"]
    cover_asset = dict(cover)
    cover_asset["docxWidthInches"] = float(cover.get("docxWidthInches", 6.7))
    add_asset_picture(document, bundle_dir, cover_asset)
    document.add_paragraph(meta["title"], style=TITLE_STYLE)
    add_article(document, extract_article_lines(article_text), bundle_dir, assets, number_id)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    font_dir = Path(__file__).resolve().parent.parent / "assets" / "fonts"
    embed_font_family(
        output,
        BODY_FONT,
        font_dir / "Poppins-Regular.ttf",
        font_dir / "Poppins-Bold.ttf",
        "origin-sculpture-poppins",
    )
    embed_font_family(
        output,
        HEADING_FONT,
        font_dir / "LibreBaskerville-Regular.ttf",
        font_dir / "LibreBaskerville-Bold.ttf",
        "origin-sculpture-libre-baskerville",
    )
    from verify_publish_docx import verify_publish_docx

    verify_publish_docx(output, bundle_dir)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("bundle_dir", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.bundle_dir.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
