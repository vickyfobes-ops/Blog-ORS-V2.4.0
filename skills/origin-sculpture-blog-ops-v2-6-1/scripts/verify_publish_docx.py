#!/usr/bin/env python3
"""Fail closed when an Origin final-upload DOCX drifts from the approved format."""

from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree as ET


BLACK = "000000"
EXPECTED_STYLES = {
    "Origin Body": ("Poppins", 12.5),
    "Origin Title": ("Libre Baskerville", 33.0),
    "Origin H2": ("Libre Baskerville", 27.0),
    "Origin H3": ("Libre Baskerville", 16.5),
    "Origin Bullet": ("Poppins", 12.5),
    "Origin Number": ("Poppins", 12.5),
    "Origin Publish Table": ("Poppins", 10.5),
}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def fail(errors: list[str], message: str) -> None:
    errors.append(message)


def visible_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() or paragraph._p.xpath(".//w:drawing"):
            yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip() or paragraph._p.xpath(".//w:drawing"):
                        yield paragraph


def verify_publish_docx(docx_path: Path, bundle_dir: Path | None = None) -> None:
    docx_path = docx_path.resolve()
    bundle_dir = bundle_dir.resolve() if bundle_dir else None
    errors: list[str] = []
    document = Document(docx_path)

    if len(document.sections) != 1:
        fail(errors, "DOCX must contain exactly one section")
    section = document.sections[0]
    geometry = {
        "page width": (section.page_width.inches, 8.5),
        "page height": (section.page_height.inches, 11.0),
        "top margin": (section.top_margin.inches, 0.72),
        "bottom margin": (section.bottom_margin.inches, 0.72),
        "left margin": (section.left_margin.inches, 0.90),
        "right margin": (section.right_margin.inches, 0.90),
    }
    for label, (actual, expected) in geometry.items():
        if abs(actual - expected) > 0.01:
            fail(errors, f"{label} is {actual:.2f} in; expected {expected:.2f} in")

    for style_name, (font_name, size_pt) in EXPECTED_STYLES.items():
        try:
            style = document.styles[style_name]
        except KeyError:
            fail(errors, f"required style is missing: {style_name}")
            continue
        if style.font.name != font_name:
            fail(errors, f"{style_name} font is {style.font.name!r}; expected {font_name!r}")
        if style.font.size is None or abs(style.font.size.pt - size_pt) > 0.05:
            actual_size = None if style.font.size is None else style.font.size.pt
            fail(errors, f"{style_name} size is {actual_size}; expected {size_pt} pt")
        if style.font.color.rgb is None or str(style.font.color.rgb).upper() != BLACK:
            fail(errors, f"{style_name} visible text color must be #{BLACK}")

    content = list(visible_paragraphs(document))
    if len(content) < 3:
        fail(errors, "DOCX needs cover, title, and article body")
    else:
        if not content[0]._p.xpath(".//w:drawing"):
            fail(errors, "the first visible block must be the cover image")
        if content[1].style.name != "Origin Title":
            fail(errors, "the public title must immediately follow the cover using Origin Title")
        if content[1]._p.xpath("./w:pPr/w:pBdr"):
            fail(errors, "the public title must not contain a border or decorative rule")

    title_paragraphs = [p for p in content if p.style.name == "Origin Title"]
    if len(title_paragraphs) != 1:
        fail(errors, f"expected exactly one Origin Title paragraph; found {len(title_paragraphs)}")

    for paragraph in content:
        if paragraph.style.name not in EXPECTED_STYLES:
            fail(errors, f"visible paragraph uses unauthorized style: {paragraph.style.name}")

    for section in document.sections:
        for story_name, story in (("header", section.header), ("footer", section.footer)):
            if any(p.text.strip() or p._p.xpath(".//w:drawing") for p in story.paragraphs):
                fail(errors, f"visible {story_name} content is not allowed")

    with zipfile.ZipFile(docx_path, "r") as archive:
        names = set(archive.namelist())
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        for color in document_xml.xpath("//w:color", namespaces={"w": W_NS}):
            value = (color.get(f"{{{W_NS}}}val") or "").upper()
            theme = color.get(f"{{{W_NS}}}themeColor")
            if theme or value != BLACK:
                fail(errors, f"non-black or theme-based visible color found: {value or theme}")
        for hyperlink in document_xml.xpath("//w:hyperlink", namespaces={"w": W_NS}):
            colors = hyperlink.xpath(".//w:color/@w:val", namespaces={"w": W_NS})
            underlines = hyperlink.xpath(".//w:u/@w:val", namespaces={"w": W_NS})
            if not colors or any(value.upper() != BLACK for value in colors):
                fail(errors, "every hyperlink must be explicitly black")
            if not underlines or any(value != "single" for value in underlines):
                fail(errors, "every hyperlink must use a single underline")

        font_parts = [name for name in names if name.startswith("word/fonts/") and name.endswith(".odttf")]
        if len(font_parts) < 4:
            fail(errors, "embedded Poppins and Libre Baskerville regular/bold font parts are missing")
        font_table = ET.fromstring(archive.read("word/fontTable.xml"))
        for family in ("Poppins", "Libre Baskerville"):
            declared = font_table.xpath(f"//w:font[@w:name='{family}']", namespaces={"w": W_NS})
            if not declared:
                fail(errors, f"{family} font-table declaration is missing")
                continue
            regular = declared[0].xpath("./w:embedRegular", namespaces={"w": W_NS})
            bold = declared[0].xpath("./w:embedBold", namespaces={"w": W_NS})
            if not regular or not bold:
                fail(errors, f"{family} regular and bold embedding declarations are required")

    alt_nodes = document.element.xpath(".//wp:docPr")
    if any(not (node.get("descr") or "").strip() for node in alt_nodes):
        fail(errors, "every embedded image must have descriptive alt text")

    if bundle_dir:
        meta = json.loads((bundle_dir / "meta.json").read_text(encoding="utf-8"))
        assets = json.loads((bundle_dir / "image-assets.json").read_text(encoding="utf-8"))
        if title_paragraphs and title_paragraphs[0].text.strip() != str(meta["title"]).strip():
            fail(errors, "the visible title does not match meta.json")
        if len(document.inline_shapes) != len(assets):
            fail(errors, f"embedded image count is {len(document.inline_shapes)}; expected {len(assets)}")
        assets_by_alt = {str(asset["alt"]): asset for asset in assets}
        for node in alt_nodes:
            alt = (node.get("descr") or "").strip()
            asset = assets_by_alt.get(alt)
            if not asset:
                fail(errors, f"embedded image alt text is not declared in image-assets.json: {alt}")
                continue
            default_width = 6.7 if asset.get("slot") == "cover" else 4.5 if asset.get("slot") == "project-review" else 5.5
            expected_width = float(asset.get("docxWidthInches", default_width))
            extents = node.xpath("../wp:extent/@cx")
            if not extents:
                fail(errors, f"embedded image width is missing for slot: {asset.get('slot', '')}")
                continue
            actual_width = int(extents[0]) / 914400
            if abs(actual_width - expected_width) > 0.02:
                fail(
                    errors,
                    f"{asset.get('slot', '')} image width is {actual_width:.2f} in; expected {expected_width:.2f} in",
                )
            if asset.get("slot") == "project-review" and not node.xpath("ancestor::w:p[1]/w:pPr/w:keepNext"):
                fail(errors, "project-review image must stay with the closing CTA heading")

    if errors:
        joined = "\n- ".join(errors)
        raise ValueError(f"Origin final-upload DOCX format check failed:\n- {joined}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--bundle", type=Path)
    args = parser.parse_args()
    verify_publish_docx(args.docx, args.bundle)
    print("Origin final-upload DOCX format check passed")


if __name__ == "__main__":
    main()
